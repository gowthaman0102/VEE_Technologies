from __future__ import annotations

import copy
import csv
import io
from xml.sax.saxutils import escape
from collections import Counter, defaultdict
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from app.ingestion.sources import get_enabled_sources

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.alert import Alert
from app.models.article import Article
from app.models.article_business_impact import ArticleBusinessImpact
from app.models.article_competitor_mention import ArticleCompetitorMention
from app.models.article_sentiment import ArticleSentiment
from app.models.article_triage import ArticleTriage
from app.models.company import Company
from app.models.company_relationship import CompanyRelationship
from app.models.event_cluster import (
    EventCluster,
    EventClusterMembership,
)
from app.models.generated_report import GeneratedReport
from app.models.risk_assessment import RiskAssessment
from app.models.risk_insight import RiskInsight
REPORT_DISPLAY_TIMEZONE = ZoneInfo(
    "Asia/Kolkata"
)


def _report_display_time(
    value: datetime,
) -> datetime:
    if value.tzinfo is None:
        value = value.replace(
            tzinfo=timezone.utc
        )

    return value.astimezone(
        REPORT_DISPLAY_TIMEZONE
    )


from app.services.analytics_service import (
    get_business_impact_distribution,
    get_event_summary,
    get_risk_summary,
    get_sentiment_distribution,
    validate_time_window,
)
from app.utils.article_metadata import publisher_name


REPORT_FORMATS = {"pdf", "xlsx", "csv"}

REPORT_CONTENT_TYPES = {
    "pdf": "application/pdf",
    "xlsx": (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),
    "csv": "text/csv",
}

REPORT_EXTENSIONS = {
    "pdf": "pdf",
    "xlsx": "xlsx",
    "csv": "csv",
}


def _article_timestamp(article: Article) -> datetime:
    return article.published_at or article.collected_at


def _iso(value: datetime | None) -> str | None:
    return value.isoformat() if value is not None else None


async def _get_company(
    db: AsyncSession,
    company_id: int,
) -> Company:
    company = await db.get(Company, company_id)

    if company is None:
        raise ValueError(
            f"Company {company_id} was not found."
        )

    return company


async def _get_report_articles(
    db: AsyncSession,
    *,
    company_id: int,
    start_date: datetime,
    end_date: datetime,
    snapshot_at: datetime,
    time_mode: str = "media",
    article_ids: list[int] | None = None,
) -> list[Article]:
    enabled_source_names = [
        source.name
        for source in get_enabled_sources()
    ]

    if time_mode == "ingestion":
        article_time = Article.collected_at
    else:
        article_time = func.coalesce(
            Article.published_at,
            Article.collected_at,
        )

    conditions = [Article.collected_at <= snapshot_at]
    if article_ids is None:
        conditions.extend([
            Article.source_name.in_(enabled_source_names),
            article_time >= start_date,
            article_time < end_date,
        ])
    else:
        conditions.append(Article.id.in_(article_ids))

    stmt = (
        select(Article)
        .where(*conditions)
        .order_by(
            article_time.desc(),
            Article.id.desc(),
        )
    )

    raw_articles = list(
        (await db.execute(stmt)).scalars().all()
    )

    unique_articles = []
    seen_external_ids = set()
    seen_urls = set()

    for article in raw_articles:
        if article.external_id and article.external_id in seen_external_ids:
            continue
            
        if getattr(article, "canonical_url", None) and article.canonical_url in seen_urls:
            continue
            
        if article.url and article.url in seen_urls:
            continue

        if article.external_id:
            seen_external_ids.add(article.external_id)

        if getattr(article, "canonical_url", None):
            seen_urls.add(article.canonical_url)

        if article.url:
            seen_urls.add(article.url)

        unique_articles.append(article)

    return unique_articles


async def _fetch_by_article_ids(
    db: AsyncSession,
    model,
    *,
    company_id: int,
    article_ids: list[int],
):
    if not article_ids:
        return []

    stmt = select(model).where(
        model.company_id == company_id,
        model.article_id.in_(article_ids),
    )

    return list(
        (
            await db.execute(stmt)
        ).scalars().all()
    )


async def _get_event_memberships(
    db: AsyncSession,
    *,
    company_id: int,
    article_ids: list[int],
) -> tuple[
    dict[int, EventClusterMembership],
    dict[int, EventCluster],
]:
    if not article_ids:
        return {}, {}

    memberships = list(
        (
            await db.execute(
                select(EventClusterMembership).where(
                    EventClusterMembership.company_id
                    == company_id,
                    EventClusterMembership.article_id.in_(
                        article_ids
                    ),
                )
            )
        ).scalars().all()
    )

    cluster_ids = {
        item.cluster_id
        for item in memberships
    }

    clusters: list[EventCluster] = []

    if cluster_ids:
        clusters = list(
            (
                await db.execute(
                    select(EventCluster).where(
                        EventCluster.company_id
                        == company_id,
                        EventCluster.id.in_(cluster_ids),
                    )
                )
            ).scalars().all()
        )

    return (
        {
            item.article_id: item
            for item in memberships
        },
        {
            item.id: item
            for item in clusters
        },
    )


async def _configured_competitor_names(
    db: AsyncSession,
    company_id: int,
) -> list[str]:
    stmt = (
        select(
            CompanyRelationship.related_company_name
        )
        .where(
            CompanyRelationship.company_id == company_id,
            CompanyRelationship.relationship_type
            == "competitor",
        )
        .order_by(
            CompanyRelationship.related_company_name.asc()
        )
    )

    return list(
        (
            await db.execute(stmt)
        ).scalars().all()
    )


def _index_by_article(rows) -> dict[int, object]:
    return {
        row.article_id: row
        for row in rows
    }


def _alerts_by_article(
    rows: list[Alert],
) -> dict[int, list[Alert]]:
    result: dict[int, list[Alert]] = defaultdict(list)

    for row in rows:
        result[row.article_id].append(row)

    return result


def _article_trend(
    articles: list[Article],
) -> list[dict]:
    counts: Counter[str] = Counter()

    for article in articles:
        period = _article_timestamp(
            article
        ).date().isoformat()

        counts[period] += 1

    return [
        {
            "period": period,
            "count": count,
        }
        for period, count in sorted(counts.items())
    ]


def _source_summary(
    article_rows: list[dict],
) -> list[dict]:
    counts: Counter[str] = Counter(
        row["source_name"]
        for row in article_rows
    )

    return [
        {
            "source_name": source,
            "article_count": count,
        }
        for source, count in counts.most_common()
    ]


def _report_summaries(
    article_rows: list[dict],
) -> tuple[dict, dict, dict, dict]:
    sentiments = Counter(
        row["sentiment"]
        for row in article_rows
        if row["sentiment"]
    )
    risk_rows = [
        row for row in article_rows
        if row["risk_score"] is not None
    ]
    risk_scores = [row["risk_score"] for row in risk_rows]
    risk = {
        "average_risk_score": round(sum(risk_scores) / len(risk_scores), 2) if risk_scores else 0.0,
        "highest_risk_score": max(risk_scores) if risk_scores else 0.0,
        "high_risk_count": sum(row["risk_level"] == "high" for row in risk_rows),
        "medium_risk_count": sum(row["risk_level"] == "medium" for row in risk_rows),
        "low_risk_count": sum(row["risk_level"] == "low" for row in risk_rows),
        "critical_risk_count": sum(row["risk_level"] == "critical" for row in risk_rows),
    }
    business = Counter()
    primary = Counter()
    for row in article_rows:
        if row["business_impact_primary"]:
            primary[row["business_impact_primary"]] += 1
        for category in row["business_impact_categories"]:
            business[category] += 1
            
    events_map = defaultdict(lambda: {"count": 0, "title": ""})
    for row in article_rows:
        if row["event_cluster_id"] is not None:
            cid = row["event_cluster_id"]
            events_map[cid]["count"] += 1
            events_map[cid]["title"] = row.get("event_cluster_title") or f"Event {cid}"
            
    largest_events = []
    for cid, info in sorted(events_map.items(), key=lambda x: x[1]["count"], reverse=True)[:5]:
        largest_events.append({
            "cluster_id": cid,
            "title": info["title"],
            "article_count": info["count"],
        })

    return (
        {"positive": sentiments["positive"], "neutral": sentiments["neutral"], "negative": sentiments["negative"]},
        risk,
        {"items": dict(business), "primary_distribution": dict(primary), "category_distribution": dict(business)},
        {"total_events": len(events_map), "largest_events": largest_events},
    )


def _competitor_summary(
    configured_names: list[str],
    article_rows: list[dict],
) -> list[dict]:
    result = []

    for name in configured_names:
        matching = [
            row
            for row in article_rows
            if name in row["competitors"]
        ]

        source_counts = Counter(
            row["source_name"]
            for row in matching
        )

        sentiment_counts = Counter(
            row["sentiment"]
            for row in matching
            if row["sentiment"]
        )

        unique_events = {
            row["event_cluster_id"]
            for row in matching
            if row["event_cluster_id"] is not None
        }

        result.append(
            {
                "name": name,
                "mention_count": len(matching),
                "article_count": len(matching),
                "event_count": len(unique_events),
                "sentiment": {
                    "positive": sentiment_counts.get(
                        "positive",
                        0,
                    ),
                    "neutral": sentiment_counts.get(
                        "neutral",
                        0,
                    ),
                    "negative": sentiment_counts.get(
                        "negative",
                        0,
                    ),
                },
                "sources": [
                    {
                        "source_name": source,
                        "count": count,
                    }
                    for source, count
                    in source_counts.most_common()
                ],
            }
        )

    return result



def build_deterministic_executive_summary(
    report_data: dict,
) -> str:
    company_name = report_data.get(
        "company_name",
        "the company",
    )

    total_articles = int(
        report_data.get(
            "total_articles",
            0,
        )
    )

    total_events = int(
        report_data.get(
            "total_events",
            0,
        )
    )

    high_risk_count = int(
        report_data.get(
            "high_risk_count",
            0,
        )
    )

    sentiment = report_data.get(
        "sentiment_balance",
        {},
    )

    positive = int(
        sentiment.get(
            "positive",
            0,
        )
    )

    neutral = int(
        sentiment.get(
            "neutral",
            0,
        )
    )

    negative = int(
        sentiment.get(
            "negative",
            0,
        )
    )

    alerts = report_data.get(
        "alerts",
        [],
    )

    configured_competitors = (
        report_data.get(
            "competitors",
            [],
        )
    )

    if total_articles == 0:
        return (
            f"No qualifying media articles were found for "
            f"{company_name} during the selected reporting "
            f"period. As a result, no media-driven sentiment, "
            f"risk, business-impact, event, competitor, or "
            f"alert conclusions are reported for this period."
        )

    sentiment_total = (
        positive
        + neutral
        + negative
    )

    if sentiment_total == 0:
        sentiment_text = (
            "No sentiment classifications were available."
        )
    else:
        dominant_label, dominant_count = max(
            (
                ("positive", positive),
                ("neutral", neutral),
                ("negative", negative),
            ),
            key=lambda item: item[1],
        )

        sentiment_text = (
            f"Media tone was predominantly "
            f"{dominant_label}, with "
            f"{dominant_count} of "
            f"{sentiment_total} classified articles."
        )

    if high_risk_count > 0:
        risk_text = (
            f"{high_risk_count} article"
            f"{'s' if high_risk_count != 1 else ''} "
            f"{'was' if high_risk_count == 1 else 'were'} "
            f"classified as high risk."
        )
    else:
        risk_text = (
            "No high-risk articles were identified."
        )

    if total_events > 0:
        event_text = (
            f"The coverage was grouped into "
            f"{total_events} tracked event"
            f"{'s' if total_events != 1 else ''}."
        )
    else:
        event_text = (
            "No qualifying event clusters were identified."
        )

    if alerts:
        alert_text = (
            f"{len(alerts)} alert"
            f"{'s' if len(alerts) != 1 else ''} "
            f"were generated from the monitored coverage."
        )
    else:
        alert_text = (
            "No alerts were generated during the period."
        )

    if configured_competitors:
        competitor_mentions = sum(
            int(
                item.get(
                    "mention_count",
                    0,
                )
            )
            for item in configured_competitors
        )

        competitor_text = (
            f"Configured competitors accounted for "
            f"{competitor_mentions} detected mention"
            f"{'s' if competitor_mentions != 1 else ''}."
        )
    else:
        competitor_text = (
            "No competitors are currently configured "
            "for comparison."
        )

    return " ".join(
        (
            f"{company_name} had "
            f"{total_articles} qualifying media article"
            f"{'s' if total_articles != 1 else ''} "
            f"during the selected reporting period.",
            sentiment_text,
            risk_text,
            event_text,
            alert_text,
            competitor_text,
        )
    )


def _highest_risk_stories(
    article_rows: list[dict],
    *,
    limit: int = 10,
) -> list[dict]:
    ranked = [
        row
        for row in article_rows
        if row["risk_score"] is not None
    ]

    ranked.sort(
        key=lambda item: (
            item["risk_score"],
            item["published_at"] or "",
        ),
        reverse=True,
    )

    return ranked[:limit]


async def build_company_report(
    db: AsyncSession,
    *,
    company_id: int,
    start_date: datetime,
    end_date: datetime,
    snapshot_at: datetime,
    time_mode: str = "media",
    article_ids: list[int] | None = None,
    report_scope: str = "standard",
    report_title: str | None = None,
) -> dict:
    start_date, end_date = validate_time_window(
        start_date,
        end_date,
    )

    company = await _get_company(
        db,
        company_id,
    )

    articles = await _get_report_articles(
        db,
        company_id=company_id,
        start_date=start_date,
        end_date=end_date,
        snapshot_at=snapshot_at,
        time_mode=time_mode,
        article_ids=article_ids,
    )

    article_ids = [
        article.id
        for article in articles
    ]

    triages = await _fetch_by_article_ids(
        db,
        ArticleTriage,
        company_id=company_id,
        article_ids=article_ids,
    )

    sentiments = await _fetch_by_article_ids(
        db,
        ArticleSentiment,
        company_id=company_id,
        article_ids=article_ids,
    )

    impacts = await _fetch_by_article_ids(
        db,
        ArticleBusinessImpact,
        company_id=company_id,
        article_ids=article_ids,
    )

    competitor_mentions = await _fetch_by_article_ids(
        db,
        ArticleCompetitorMention,
        company_id=company_id,
        article_ids=article_ids,
    )

    risks = await _fetch_by_article_ids(
        db,
        RiskAssessment,
        company_id=company_id,
        article_ids=article_ids,
    )

    risk_insights = await _fetch_by_article_ids(
        db,
        RiskInsight,
        company_id=company_id,
        article_ids=article_ids,
    )

    alerts = await _fetch_by_article_ids(
        db,
        Alert,
        company_id=company_id,
        article_ids=article_ids,
    )

    membership_by_article, cluster_by_id = (
        await _get_event_memberships(
            db,
            company_id=company_id,
            article_ids=article_ids,
        )
    )

    triage_by_article = _index_by_article(
        triages
    )
    sentiment_by_article = _index_by_article(
        sentiments
    )
    impact_by_article = _index_by_article(
        impacts
    )
    competitors_by_article = _index_by_article(
        competitor_mentions
    )
    risk_by_article = _index_by_article(
        risks
    )
    insight_by_article = _index_by_article(
        risk_insights
    )
    alerts_by_article = _alerts_by_article(
        alerts
    )

    article_rows: list[dict] = []

    for article in articles:
        triage = triage_by_article.get(
            article.id
        )
        sentiment = sentiment_by_article.get(
            article.id
        )
        impact = impact_by_article.get(
            article.id
        )
        competitor = competitors_by_article.get(
            article.id
        )
        risk = risk_by_article.get(
            article.id
        )
        insight = insight_by_article.get(
            article.id
        )
        membership = membership_by_article.get(
            article.id
        )

        cluster = (
            cluster_by_id.get(
                membership.cluster_id
            )
            if membership is not None
            else None
        )

        article_alerts = alerts_by_article.get(
            article.id,
            [],
        )

        article_rows.append(
            {
                "article_id": article.id,
                "title": article.title,
                "publisher_name": publisher_name(
                    article.source_name,
                    article.title,
                    article.url,
                    getattr(article, "canonical_url", None),
                ),
                "source_name": article.source_name,
                "url": article.url,
                "has_triage": triage is not None,
                "published_at": _iso(article.published_at),
                "collected_at": _iso(
                    getattr(
                        article,
                        "collected_at",
                        getattr(article, "created_at", None),
                    )
                ),
                "sentiment": (
                    sentiment.label
                    if sentiment is not None
                    else None
                ),
                "sentiment_score": (
                    float(sentiment.score)
                    if sentiment is not None
                    else None
                ),
                "business_impact_primary": (
                    impact.primary_category
                    if impact is not None
                    else None
                ),
                "business_impact_categories": (
                    list(impact.categories or [])
                    if impact is not None
                    else []
                ),
                "business_impact_summary": (
                    impact.impact_summary
                    if impact is not None
                    else None
                ),
                "competitors": (
                    list(
                        competitor.competitors
                        or []
                    )
                    if competitor is not None
                    else []
                ),
                "risk_score": (
                    float(risk.risk_score)
                    if risk is not None
                    else None
                ),
                "risk_level": (
                    risk.risk_level
                    if risk is not None
                    else None
                ),
                "escalation_action": (
                    risk.escalation_action
                    if risk is not None
                    else None
                ),
                "attention_level": (
                    insight.attention_level
                    if insight is not None
                    else None
                ),
                "risk_headline": (
                    insight.headline
                    if insight is not None
                    else None
                ),
                "risk_summary": (
                    insight.executive_summary
                    if insight is not None
                    else None
                ),
                "event_cluster_id": (
                    cluster.id
                    if cluster is not None
                    else None
                ),
                "event_cluster_title": (
                    cluster.title
                    if cluster is not None
                    else None
                ),
                "alerts": [
                    {
                        "alert_type": item.alert_type,
                        "severity": item.severity,
                        "title": item.title,
                        "delivery_status": (
                            item.delivery_status
                        ),
                        "created_at": _iso(
                            item.created_at
                        ),
                    }
                    for item in article_alerts
                ],
            }
        )

    (
        sentiment_summary,
        risk_summary,
        business_summary,
        event_summary,
    ) = _report_summaries(article_rows)

    configured_competitors = (
        await _configured_competitor_names(
            db,
            company_id,
        )
    )

    sources = _source_summary(
        article_rows
    )

    competitors = _competitor_summary(
        configured_competitors,
        article_rows,
    )

    alert_rows = [
        {
            "article_id": alert.article_id,
            "alert_type": alert.alert_type,
            "severity": alert.severity,
            "title": alert.title,
            "message": alert.message,
            "delivery_status": (
                alert.delivery_status
            ),
            "delivery_channel": (
                alert.delivery_channel
            ),
            "sla_due_at": _iso(
                alert.sla_due_at
            ),
            "delivered_at": _iso(
                alert.delivered_at
            ),
            "created_at": _iso(
                alert.created_at
            ),
        }
        for alert in alerts
    ]

    metrics = [
        {
            "label": "Total articles",
            "value": len(article_rows),
        },
        {
            "label": "Total events",
            "value": event_summary.get(
                "total_events",
                0,
            ),
        },
        {
            "label": "Average risk score",
            "value": risk_summary.get(
                "average_risk_score",
                0.0,
            ),
        },
        {
            "label": "Highest risk score",
            "value": risk_summary.get(
                "highest_risk_score",
                0.0,
            ),
        },
        {
            "label": "High risk count",
            "value": risk_summary.get(
                "high_risk_count",
                0,
            ),
        },
        {
            "label": "Positive sentiment",
            "value": sentiment_summary.get(
                "positive",
                0,
            ),
        },
        {
            "label": "Neutral sentiment",
            "value": sentiment_summary.get(
                "neutral",
                0,
            ),
        },
        {
            "label": "Negative sentiment",
            "value": sentiment_summary.get(
                "negative",
                0,
            ),
        },
        {
            "label": "Alerts",
            "value": len(alert_rows),
        },
    ]

    processed_count = sum(1 for row in article_rows if row.get("has_triage"))
    sentiment_count = sum(1 for row in article_rows if row.get("sentiment") is not None)
    risk_count = sum(1 for row in article_rows if row.get("risk_score") is not None)
    impact_count = sum(1 for row in article_rows if row.get("business_impact_primary") is not None)
    event_assigned_count = sum(1 for row in article_rows if row.get("event_cluster_id") is not None)

    metrics.extend([
        {"label": "Processed Intelligence", "value": processed_count},
        {"label": "Sentiment Analyzed", "value": sentiment_count},
        {"label": "Risk Assessed", "value": risk_count},
        {"label": "Business Impact Analyzed", "value": impact_count},
        {"label": "Event Assigned", "value": event_assigned_count},
    ])

    report_data = {
        "company_id": company_id,
        "company_name": company.name,
        "start_date": start_date,
        "end_date": end_date,
        "snapshot_at": snapshot_at,
        "time_mode": time_mode,
        "total_articles": len(article_rows),
        "total_events": event_summary.get(
            "total_events",
            0,
        ),
        "high_risk_count": risk_summary.get(
            "high_risk_count",
            0,
        ),
        "medium_risk_count": risk_summary.get(
            "medium_risk_count",
            0,
        ),
        "low_risk_count": risk_summary.get(
            "low_risk_count",
            0,
        ),
        "sentiment_balance": {
            "positive": sentiment_summary.get(
                "positive",
                0,
            ),
            "neutral": sentiment_summary.get(
                "neutral",
                0,
            ),
            "negative": sentiment_summary.get(
                "negative",
                0,
            ),
        },
        "metrics": metrics,
        "overview": {
            "total_articles": len(
                article_rows
            ),
            "total_events": event_summary.get(
                "total_events",
                0,
            ),
            "total_alerts": len(
                alert_rows
            ),
        },
        "article_trend": _article_trend(
            articles
        ),
        "articles": article_rows,
        "report_scope": report_scope,
        "report_title": report_title,
        "sentiment": sentiment_summary,
        "risk": risk_summary,
        "business_impact": business_summary,
        "events": event_summary,
        "sources": sources,
        "competitors": competitors,
        "alerts": alert_rows,
        "highest_risk_stories": (
            _highest_risk_stories(
                article_rows
            )
        ),
    }

    report_data["executive_summary"] = (
        build_deterministic_executive_summary(
            report_data
        )
    )

    return report_data


def report_rows(
    report_data: dict,
) -> list[tuple[str, str]]:
    risk = report_data.get("risk", {})
    rows = [
        (
            "company",
            report_data.get(
                "company_name",
                str(
                    report_data[
                        "company_id"
                    ]
                ),
            ),
        ),
        (
            "start_date",
            report_data[
                "start_date"
            ].isoformat(),
        ),
        (
            "end_date",
            report_data[
                "end_date"
            ].isoformat(),
        ),
        (
            "total_articles",
            str(
                report_data[
                    "total_articles"
                ]
            ),
        ),
        (
            "total_events",
            str(
                report_data[
                    "total_events"
                ]
            ),
        ),
        (
            "high_risk_count",
            str(
                report_data[
                    "high_risk_count"
                ]
            ),
        ),
        (
            "medium_risk_count",
            str(
                report_data[
                    "medium_risk_count"
                ]
            ),
        ),
        (
            "low_risk_count",
            str(
                report_data[
                    "low_risk_count"
                ]
            ),
        ),
        (
            "positive_sentiment",
            str(
                report_data[
                    "sentiment_balance"
                ]["positive"]
            ),
        ),
        (
            "neutral_sentiment",
            str(
                report_data[
                    "sentiment_balance"
                ]["neutral"]
            ),
        ),
        (
            "negative_sentiment",
            str(
                report_data[
                    "sentiment_balance"
                ]["negative"]
            ),
        ),
        (
            "average_risk_score",
            str(
                risk.get(
                    "average_risk_score",
                    report_data.get("average_risk_score", 0.0),
                )
            ),
        ),
        (
            "highest_risk_score",
            str(
                risk.get(
                    "highest_risk_score",
                    report_data.get("highest_risk_score", 0.0),
                )
            ),
        ),
        (
            "source_count",
            str(len(report_data.get("sources", []))),
        ),
        (
            "alert_count",
            str(len(report_data.get("alerts", []))),
        ),
    ]

    return rows


def export_report_csv(
    report_data: dict,
) -> bytes:
    output = io.StringIO(
        newline=""
    )
    writer = csv.writer(
        output
    )

    writer.writerow(
        (
            "row_type",
            "metric_label",
            "article_id",
            "title",
            "publisher",
            "source_provider",
            "url",
            "published_at",
            "collected_at",
            "sentiment",
            "risk_level",
            "risk_score",
            "business_impact",
            "event_cluster_id",
            "event_cluster_title",
        )
    )

    for metric, value in report_rows(report_data):
        writer.writerow(
            (
                "metric",
                metric,
                value,
                "", "", "", "", "", "", "", "", "", "", "", "",
            )
        )

    articles = report_data.get("articles", [])
    for item in articles:
        writer.writerow(
            (
                "article",
                "",
                item.get("article_id"),
                item.get("title"),
                item.get("publisher_name"),
                item.get("source_name"),
                item.get("url"),
                item.get("published_at"),
                item.get("collected_at"),
                item.get("sentiment"),
                item.get("risk_level"),
                item.get("risk_score"),
                item.get("business_impact_primary"),
                item.get("event_cluster_id"),
                item.get("event_cluster_title"),
            )
        )

    return output.getvalue().encode(
        "utf-8"
    )


def _is_article_collection_report(report_data: dict) -> bool:
    return report_data.get("report_scope") in {"search", "analytics"}


def export_article_collection_csv(report_data: dict) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(("report", report_data.get("report_title") or "Article report"))
    writer.writerow(("title", "publisher", "published_at", "sentiment", "critical_level"))
    for item in report_data.get("articles", []):
        writer.writerow((
            item.get("title"),
            item.get("publisher_name"),
            item.get("published_at") or item.get("collected_at"),
            item.get("sentiment"),
            item.get("risk_level"),
        ))
    return output.getvalue().encode("utf-8")


def export_article_collection_xlsx(report_data: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Articles"
    sheet.append([report_data.get("report_title") or "Article report"])
    sheet.append(["Title", "Publisher", "Published At", "Sentiment", "Critical Level"])
    for item in report_data.get("articles", []):
        sheet.append([
            item.get("title"),
            item.get("publisher_name"),
            item.get("published_at") or item.get("collected_at"),
            item.get("sentiment"),
            item.get("risk_level"),
        ])
    for cell in sheet[1] + sheet[2]:
        _xlsx_header_style(cell)
    for column, width in zip("ABCDE", (70, 28, 24, 16, 18)):
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_article_collection_pdf(report_data: dict) -> bytes:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    title_style = ParagraphStyle(
        "ArticleReportTitle",
        parent=styles["Title"],
        alignment=TA_LEFT,
        spaceAfter=12,
    )
    header_style = ParagraphStyle(
        "ArticleReportHeader",
        parent=styles["BodyText"],
        alignment=TA_CENTER,
        textColor=colors.white,
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
    )
    body_style = ParagraphStyle(
        "ArticleReportBody",
        parent=styles["BodyText"],
        alignment=TA_LEFT,
        fontSize=8,
        leading=10,
        spaceAfter=0,
    )
    centered_body_style = ParagraphStyle(
        "ArticleReportCenteredBody",
        parent=body_style,
        alignment=TA_CENTER,
    )
    date_body_style = ParagraphStyle(
        "ArticleReportDateBody",
        parent=centered_body_style,
        leading=9,
    )

    story = [
        Paragraph(
            escape(report_data.get("report_title") or "Article report"),
            title_style,
        ),
    ]
    rows = [[
        Paragraph("Title", header_style),
        Paragraph("Publisher", header_style),
        Paragraph("Published", header_style),
        Paragraph("Sentiment", header_style),
        Paragraph("Critical level", header_style),
    ]]
    for item in report_data.get("articles", []):
        published_at = item.get("published_at") or item.get("collected_at")
        rows.append([
            Paragraph(escape(str(item.get("title") or "—")), body_style),
            Paragraph(escape(str(item.get("publisher_name") or "—")), body_style),
            Paragraph(
                _pdf_format_timestamp(published_at),
                date_body_style,
            ),
            Paragraph(escape(str(item.get("sentiment") or "—")), centered_body_style),
            Paragraph(escape(str(item.get("risk_level") or "—")), centered_body_style),
        ])
    table = Table(
        rows,
        repeatRows=1,
        colWidths=[2.9 * inch, 1.25 * inch, 1.55 * inch, 0.8 * inch, 1.0 * inch],
        hAlign="LEFT",
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FC")]),
    ]))
    story.append(table)
    document.build(story)
    return output.getvalue()


def _pdf_format_timestamp(value: object | None) -> str:
    if value is None:
        return "Not available"

    parsed: datetime | None = None
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value).strip()
        if text:
            try:
                parsed = datetime.fromisoformat(
                    text.replace("Z", "+00:00")
                )
            except ValueError:
                return escape(text)

    if parsed is None:
        return "Not available"

    return escape(
        _report_display_time(parsed).strftime(
            "%d %b %Y\n%H:%M IST"
        )
    ).replace("\n", "<br/>")


def _xlsx_header_style(
    cell,
) -> None:
    cell.fill = PatternFill(
        "solid",
        fgColor="16324F",
    )
    cell.font = Font(
        color="FFFFFF",
        bold=True,
    )
    cell.alignment = Alignment(
        vertical="center",
        wrap_text=True,
    )

    thin = Side(
        style="thin",
        color="D1D5DB",
    )

    cell.border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )


def _xlsx_body_style(
    cell,
    *,
    wrap: bool = True,
) -> None:
    cell.alignment = Alignment(
        vertical="top",
        wrap_text=wrap,
    )

    thin = Side(
        style="thin",
        color="E5E7EB",
    )

    cell.border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )


def _xlsx_title_style(
    cell,
) -> None:
    cell.font = Font(
        size=18,
        bold=True,
        color="16324F",
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


def _xlsx_section_style(
    cell,
) -> None:
    cell.font = Font(
        size=12,
        bold=True,
        color="16324F",
    )

    cell.fill = PatternFill(
        "solid",
        fgColor="EAF0F6",
    )

    cell.alignment = Alignment(
        vertical="center",
    )


def _xlsx_format_timestamp(
    value: object | None,
) -> str:
    if isinstance(value, datetime):
        return value.strftime(
            "%d %b %Y %H:%M UTC"
        )

    if value is None:
        return "Not available"

    text = str(value).strip()

    return text or "Not available"


def _xlsx_add_table(
    sheet,
    headers: list[str],
    rows: list[list],
    *,
    widths: list[float],
    empty_message: str | None = None,
) -> None:
    sheet.append(headers)

    for cell in sheet[1]:
        _xlsx_header_style(cell)

    if rows:
        for row in rows:
            sheet.append(row)
    elif empty_message:
        sheet.append(
            [
                empty_message,
                *(
                    ""
                    for _ in range(
                        max(
                            len(headers) - 1,
                            0,
                        )
                    )
                ),
            ]
        )

    for row in sheet.iter_rows(
        min_row=2,
        max_row=sheet.max_row,
    ):
        for cell in row:
            _xlsx_body_style(cell)

    sheet.freeze_panes = "A2"

    if sheet.max_row >= 1:
        sheet.auto_filter.ref = (
            f"A1:"
            f"{sheet.cell(1, len(headers)).coordinate}"
        )

    for index, width in enumerate(
        widths,
        start=1,
    ):
        sheet.column_dimensions[
            sheet.cell(
                1,
                index,
            ).column_letter
        ].width = width


def export_report_xlsx(
    report_data: dict,
) -> bytes:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"

    for name in (
        "Articles",
        "Sentiment",
        "Risk",
        "Business Impact",
        "Events",
        "Sources",
        "Competitors",
        "Alerts",
    ):
        workbook.create_sheet(
            title=name
        )

    company_name = report_data.get(
        "company_name",
        "Monitored Company",
    )

    start_date = report_data.get(
        "start_date"
    )
    end_date = report_data.get(
        "end_date"
    )

    executive_summary = report_data.get(
        "executive_summary",
        "No executive summary is available.",
    )

    summary.merge_cells(
        "A1:D1"
    )
    summary["A1"] = (
        f"{company_name} "
        "Media Intelligence Report"
    )
    _xlsx_title_style(
        summary["A1"]
    )

    summary.merge_cells(
        "A2:D2"
    )
    summary["A2"] = company_name
    summary["A2"].font = Font(
        size=12,
        bold=True,
        color="4B5563",
    )
    summary["A2"].alignment = Alignment(
        horizontal="center",
    )

    summary["A4"] = "Reporting Period"
    summary["B4"] = (
        f"{_xlsx_format_timestamp(start_date)}"
        " - "
        f"{_xlsx_format_timestamp(end_date)}"
    )

    summary["A5"] = "Generated"
    summary["B5"] = _xlsx_format_timestamp(
        datetime.now(timezone.utc)
    )

    summary["A7"] = "Executive Summary"
    summary.merge_cells(
        "A7:D7"
    )
    _xlsx_section_style(
        summary["A7"]
    )

    summary.merge_cells(
        "A8:D10"
    )
    summary["A8"] = executive_summary
    summary["A8"].alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    summary["A12"] = "KPI Summary"
    summary.merge_cells(
        "A12:D12"
    )
    _xlsx_section_style(
        summary["A12"]
    )

    sentiment_balance = report_data.get(
        "sentiment_balance",
        {},
    )

    risk = report_data.get(
        "risk",
        {},
    )

    kpis = [
        (
            "Total Articles",
            report_data.get(
                "total_articles",
                0,
            ),
        ),
        (
            "Total Events",
            report_data.get(
                "total_events",
                0,
            ),
        ),
        (
            "High Risk",
            report_data.get(
                "high_risk_count",
                0,
            ),
        ),
        (
            "Medium Risk",
            report_data.get(
                "medium_risk_count",
                0,
            ),
        ),
        (
            "Low Risk",
            report_data.get(
                "low_risk_count",
                0,
            ),
        ),
        (
            "Positive Sentiment",
            sentiment_balance.get(
                "positive",
                0,
            ),
        ),
        (
            "Neutral Sentiment",
            sentiment_balance.get(
                "neutral",
                0,
            ),
        ),
        (
            "Negative Sentiment",
            sentiment_balance.get(
                "negative",
                0,
            ),
        ),
        (
            "Alerts",
            len(
                report_data.get(
                    "alerts",
                    [],
                )
            ),
        ),
        (
            "Average Risk Score",
            risk.get(
                "average_risk_score",
                0.0,
            ),
        ),
        (
            "Highest Risk Score",
            risk.get(
                "highest_risk_score",
                0.0,
            ),
        ),
    ]

    summary.append(
        [
            "Metric",
            "Value",
        ]
    )

    header_row = summary.max_row

    for cell in summary[
        header_row
    ]:
        if cell.column <= 2:
            _xlsx_header_style(cell)

    for label, value in kpis:
        summary.append(
            [
                label,
                value,
            ]
        )

    for row in summary.iter_rows(
        min_row=header_row + 1,
        max_row=summary.max_row,
        min_col=1,
        max_col=2,
    ):
        for cell in row:
            _xlsx_body_style(cell)

    summary.column_dimensions[
        "A"
    ].width = 28
    summary.column_dimensions[
        "B"
    ].width = 30
    summary.column_dimensions[
        "C"
    ].width = 24
    summary.column_dimensions[
        "D"
    ].width = 24

    summary.row_dimensions[
        1
    ].height = 28

    summary.row_dimensions[
        8
    ].height = 50

    summary.sheet_view.showGridLines = False

    articles_sheet = workbook["Articles"]

    article_rows = []

    for item in report_data.get(
        "articles",
        [],
    ):
        competitors = item.get(
            "competitors",
            [],
        )

        impact_categories = item.get(
            "business_impact_categories",
            [],
        )

        nested_alerts = item.get(
            "alerts",
            [],
        )

        article_rows.append(
            [
                item.get("article_id"),
                item.get("title"),
                item.get("publisher_name"),
                item.get("source_name"),
                item.get("published_at"),
                item.get("collected_at"),
                item.get("url"),
                item.get("sentiment"),
                item.get("sentiment_score"),
                item.get("business_impact_primary"),
                ", ".join(str(value) for value in impact_categories),
                item.get("business_impact_summary"),
                ", ".join(str(value) for value in competitors),
                item.get("risk_score"),
                item.get("risk_level"),
                item.get("escalation_action"),
                item.get("attention_level"),
                item.get("risk_headline"),
                item.get("risk_summary"),
                item.get("event_cluster_title"),
                len(nested_alerts),
            ]
        )

    _xlsx_add_table(
        articles_sheet,
        [
            "Article ID",
            "Title",
            "Publisher",
            "Source Provider",
            "Published At",
            "Collected At",
            "URL",
            "Sentiment",
            "Sentiment Score",
            "Primary Impact",
            "Impact Categories",
            "Impact Summary",
            "Competitors",
            "Risk Score",
            "Risk Level",
            "Escalation",
            "Attention",
            "Risk Headline",
            "Risk Summary",
            "Event Cluster",
            "Alert Count",
        ],
        article_rows,
        widths=[
            12,
            42,
            24,
            24,
            24,
            42,
            14,
            16,
            18,
            30,
            42,
            28,
            14,
            14,
            18,
            18,
            36,
            45,
            36,
            12,
            24,
            24,
            24,
        ],
        empty_message=(
            "No qualifying articles are available "
            "for this reporting period."
        ),
    )

    articles_sheet.sheet_view.showGridLines = False

    sentiment_sheet = workbook["Sentiment"]

    sentiment = report_data.get(
        "sentiment",
        {},
    )

    sentiment_rows = [
        [
            "Positive",
            sentiment.get(
                "positive",
                0,
            ),
        ],
        [
            "Neutral",
            sentiment.get(
                "neutral",
                0,
            ),
        ],
        [
            "Negative",
            sentiment.get(
                "negative",
                0,
            ),
        ],
    ]

    _xlsx_add_table(
        sentiment_sheet,
        [
            "Classification",
            "Article Count",
        ],
        sentiment_rows,
        widths=[
            22,
            18,
        ],
    )

    sentiment_sheet[
        "D1"
    ] = "Trend"

    _xlsx_section_style(
        sentiment_sheet["D1"]
    )

    trend_headers = [
        "Period",
        "Positive",
        "Neutral",
        "Negative",
    ]

    for column, value in enumerate(
        trend_headers,
        start=4,
    ):
        cell = sentiment_sheet.cell(
            row=2,
            column=column,
            value=value,
        )
        _xlsx_header_style(cell)

    sentiment_series = sentiment.get(
        "series",
        [],
    )

    if sentiment_series:
        for row_index, item in enumerate(
            sentiment_series,
            start=3,
        ):
            values = [
                item.get("period"),
                item.get("positive", 0),
                item.get("neutral", 0),
                item.get("negative", 0),
            ]

            for column, value in enumerate(
                values,
                start=4,
            ):
                cell = sentiment_sheet.cell(
                    row=row_index,
                    column=column,
                    value=value,
                )
                _xlsx_body_style(cell)
    else:
        sentiment_sheet[
            "D3"
        ] = (
            "No sentiment trend data is "
            "available for this period."
        )
        _xlsx_body_style(
            sentiment_sheet["D3"]
        )

    sentiment_sheet.column_dimensions[
        "D"
    ].width = 24
    sentiment_sheet.column_dimensions[
        "E"
    ].width = 14
    sentiment_sheet.column_dimensions[
        "F"
    ].width = 14
    sentiment_sheet.column_dimensions[
        "G"
    ].width = 14

    sentiment_sheet.sheet_view.showGridLines = False

    risk_sheet = workbook["Risk"]

    risk = report_data.get(
        "risk",
        {},
    )

    risk_rows = [
        [
            "Average Risk Score",
            risk.get(
                "average_risk_score",
                0.0,
            ),
        ],
        [
            "Highest Risk Score",
            risk.get(
                "highest_risk_score",
                0.0,
            ),
        ],
        [
            "Low Risk",
            risk.get(
                "low_risk_count",
                0,
            ),
        ],
        [
            "Medium Risk",
            risk.get(
                "medium_risk_count",
                0,
            ),
        ],
        [
            "High Risk",
            risk.get(
                "high_risk_count",
                0,
            ),
        ],
    ]

    _xlsx_add_table(
        risk_sheet,
        [
            "Metric",
            "Value",
        ],
        risk_rows,
        widths=[
            28,
            18,
        ],
    )

    risk_sheet["D1"] = "Risk Trend"
    _xlsx_section_style(
        risk_sheet["D1"]
    )

    risk_trend_headers = [
        "Period",
        "Average Risk",
        "Maximum Risk",
        "Low",
        "Medium",
        "High",
        "Escalations",
    ]

    for column, value in enumerate(
        risk_trend_headers,
        start=4,
    ):
        cell = risk_sheet.cell(
            row=2,
            column=column,
            value=value,
        )
        _xlsx_header_style(cell)

    risk_series = risk.get(
        "series",
        [],
    )

    if risk_series:
        for row_index, item in enumerate(
            risk_series,
            start=3,
        ):
            values = [
                item.get("period"),
                item.get(
                    "average_risk_score",
                    0.0,
                ),
                item.get(
                    "maximum_risk_score",
                    0.0,
                ),
                item.get(
                    "low_count",
                    0,
                ),
                item.get(
                    "medium_count",
                    0,
                ),
                item.get(
                    "high_count",
                    0,
                ),
                item.get(
                    "escalation_count",
                    0,
                ),
            ]

            for column, value in enumerate(
                values,
                start=4,
            ):
                cell = risk_sheet.cell(
                    row=row_index,
                    column=column,
                    value=value,
                )
                _xlsx_body_style(cell)
    else:
        risk_sheet["D3"] = (
            "No risk trend data is "
            "available for this period."
        )
        _xlsx_body_style(
            risk_sheet["D3"]
        )

    for column, width in {
        "D": 22,
        "E": 18,
        "F": 18,
        "G": 12,
        "H": 12,
        "I": 12,
        "J": 16,
    }.items():
        risk_sheet.column_dimensions[
            column
        ].width = width

    risk_sheet.sheet_view.showGridLines = False

    impact_sheet = workbook[
        "Business Impact"
    ]

    business_impact = report_data.get(
        "business_impact",
        {},
    )

    primary_distribution = (
        business_impact.get(
            "primary_distribution",
            {},
        )
    )

    category_distribution = (
        business_impact.get(
            "category_distribution",
            {},
        )
    )

    primary_rows = [
        [
            str(category)
            .replace("_", " ")
            .title(),
            count,
        ]
        for category, count
        in primary_distribution.items()
        if count
    ]

    _xlsx_add_table(
        impact_sheet,
        [
            "Primary Category",
            "Article Count",
        ],
        primary_rows,
        widths=[
            28,
            18,
        ],
        empty_message=(
            "No primary business-impact "
            "classifications are available "
            "for this period."
        ),
    )

    impact_sheet["D1"] = (
        "All-Category Distribution"
    )
    _xlsx_section_style(
        impact_sheet["D1"]
    )

    impact_sheet["D2"] = (
        "Impact Category"
    )
    impact_sheet["E2"] = (
        "Article Count"
    )

    _xlsx_header_style(
        impact_sheet["D2"]
    )
    _xlsx_header_style(
        impact_sheet["E2"]
    )

    category_rows = [
        [
            str(category)
            .replace("_", " ")
            .title(),
            count,
        ]
        for category, count
        in category_distribution.items()
        if count
    ]

    if category_rows:
        for row_index, values in enumerate(
            category_rows,
            start=3,
        ):
            for column, value in enumerate(
                values,
                start=4,
            ):
                cell = impact_sheet.cell(
                    row=row_index,
                    column=column,
                    value=value,
                )
                _xlsx_body_style(cell)
    else:
        impact_sheet["D3"] = (
            "No all-category business-impact "
            "classifications are available "
            "for this period."
        )
        _xlsx_body_style(
            impact_sheet["D3"]
        )

    impact_sheet.column_dimensions[
        "D"
    ].width = 32

    impact_sheet.column_dimensions[
        "E"
    ].width = 18

    impact_sheet.sheet_view.showGridLines = False

    events_sheet = workbook["Events"]

    events = report_data.get(
        "events",
        {},
    )

    event_rows = [
        [
            item.get("cluster_id"),
            item.get("title"),
            item.get("article_count", 0),
            item.get("first_published_at"),
            item.get("last_published_at"),
        ]
        for item in events.get(
            "largest_events",
            [],
        )
    ]

    _xlsx_add_table(
        events_sheet,
        [
            "Cluster ID",
            "Event Title",
            "Article Count",
            "First Published",
            "Last Published",
        ],
        event_rows,
        widths=[
            14,
            48,
            16,
            26,
            26,
        ],
        empty_message=(
            "No qualifying event clusters "
            "were identified for this period."
        ),
    )

    events_sheet.sheet_view.showGridLines = False

    sources_sheet = workbook["Sources"]

    source_rows = [
        [
            item.get("source_name"),
            item.get("article_count", 0),
        ]
        for item in report_data.get(
            "sources",
            [],
        )
    ]

    _xlsx_add_table(
        sources_sheet,
        [
            "Source",
            "Article Count",
        ],
        source_rows,
        widths=[
            40,
            18,
        ],
        empty_message=(
            "No qualifying media sources "
            "were identified for this period."
        ),
    )

    sources_sheet.sheet_view.showGridLines = False

    competitors_sheet = workbook[
        "Competitors"
    ]

    competitor_rows = []

    for item in report_data.get(
        "competitors",
        [],
    ):
        sentiment = item.get(
            "sentiment",
            {},
        )

        source_breakdown = ", ".join(
            (
                f"{source.get('source_name')}: "
                f"{source.get('count', 0)}"
            )
            for source in item.get(
                "sources",
                [],
            )
        )

        competitor_rows.append(
            [
                item.get("name"),
                item.get(
                    "mention_count",
                    0,
                ),
                item.get(
                    "article_count",
                    0,
                ),
                item.get(
                    "event_count",
                    0,
                ),
                sentiment.get(
                    "positive",
                    0,
                ),
                sentiment.get(
                    "neutral",
                    0,
                ),
                sentiment.get(
                    "negative",
                    0,
                ),
                source_breakdown,
            ]
        )

    _xlsx_add_table(
        competitors_sheet,
        [
            "Competitor",
            "Mentions",
            "Articles",
            "Events",
            "Positive",
            "Neutral",
            "Negative",
            "Source Breakdown",
        ],
        competitor_rows,
        widths=[
            32,
            14,
            14,
            14,
            14,
            14,
            14,
            48,
        ],
        empty_message=(
            "No competitors are currently "
            "configured for comparison."
        ),
    )

    competitors_sheet.sheet_view.showGridLines = False

    alerts_sheet = workbook["Alerts"]

    alert_rows = [
        [
            item.get("article_id"),
            item.get("alert_type"),
            item.get("severity"),
            item.get("title"),
            item.get("message"),
            item.get("delivery_status"),
            item.get("delivery_channel"),
            item.get("sla_due_at"),
            item.get("delivered_at"),
            item.get("created_at"),
        ]
        for item in report_data.get(
            "alerts",
            [],
        )
    ]

    _xlsx_add_table(
        alerts_sheet,
        [
            "Article ID",
            "Alert Type",
            "Severity",
            "Title",
            "Message",
            "Delivery Status",
            "Delivery Channel",
            "SLA Due",
            "Delivered",
            "Created",
        ],
        alert_rows,
        widths=[
            14,
            22,
            14,
            36,
            55,
            20,
            20,
            26,
            26,
            26,
        ],
        empty_message=(
            "No alerts were generated "
            "during the reporting period."
        ),
    )

    alerts_sheet.sheet_view.showGridLines = False

    output = io.BytesIO()

    workbook.save(
        output
    )

    return output.getvalue()

def _pdf_safe_text(
    value: object | None,
    *,
    fallback: str = "Not available",
) -> str:
    if value is None:
        return escape(fallback)

    text = str(value).strip()

    if not text:
        return escape(fallback)

    return escape(text)


def _pdf_display_value(
    value: object | None,
    *,
    fallback: str = "Not available",
) -> str:
    if value is None:
        return fallback

    if isinstance(value, float):
        return f"{value:.2f}"

    if isinstance(value, datetime):
        return value.strftime(
            "%d %b %Y, %H:%M UTC"
        )

    text = str(value).strip()

    return text if text else fallback


def _pdf_styles() -> dict[str, ParagraphStyle]:
    sample = getSampleStyleSheet()

    return {
        "title": ParagraphStyle(
            "VeeReportTitle",
            parent=sample["Title"],
            fontName="Helvetica-Bold",
            fontSize=22,
            leading=27,
            alignment=TA_CENTER,
            spaceAfter=12,
            textColor=colors.HexColor("#16324F"),
        ),
        "subtitle": ParagraphStyle(
            "VeeReportSubtitle",
            parent=sample["Normal"],
            fontName="Helvetica",
            fontSize=11,
            leading=15,
            alignment=TA_CENTER,
            spaceAfter=6,
            textColor=colors.HexColor("#4B5563"),
        ),
        "section": ParagraphStyle(
            "VeeReportSection",
            parent=sample["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=18,
            spaceBefore=10,
            spaceAfter=8,
            textColor=colors.HexColor("#16324F"),
        ),
        "body": ParagraphStyle(
            "VeeReportBody",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            spaceAfter=7,
            textColor=colors.HexColor("#1F2937"),
        ),
        "small": ParagraphStyle(
            "VeeReportSmall",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#4B5563"),
        ),
        "table_header": ParagraphStyle(
            "VeeReportTableHeader",
            parent=sample["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=11,
            textColor=colors.white,
        ),
        "table_cell": ParagraphStyle(
            "VeeReportTableCell",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#1F2937"),
        ),
        "empty": ParagraphStyle(
            "VeeReportEmpty",
            parent=sample["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#6B7280"),
        ),
    }


def _pdf_paragraph(
    value: object | None,
    style: ParagraphStyle,
    *,
    fallback: str = "Not available",
) -> Paragraph:
    return Paragraph(
        _pdf_safe_text(
            value,
            fallback=fallback,
        ),
        style,
    )

def _pdf_page_footer(
    canvas,
    document,
) -> None:
    canvas.saveState()

    page_width, _ = letter

    canvas.setFont(
        "Helvetica",
        7.5,
    )
    canvas.setFillColor(
        colors.HexColor("#6B7280")
    )

    canvas.drawString(
        document.leftMargin,
        0.42 * inch,
        "AI Media Intelligence",
    )

    canvas.drawRightString(
        page_width - document.rightMargin,
        0.42 * inch,
        f"Page {document.page}",
    )

    canvas.restoreState()


def export_report_pdf(
    report_data: dict,
) -> bytes:
    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title=(
            "Monitored Company Media Intelligence Report"
        ),
        author="AI Media Intelligence",
    )

    styles = _pdf_styles()
    story = []

    company_name = report_data.get(
        "company_name",
        "Monitored Company",
    )

    start_date = report_data.get(
        "start_date"
    )
    end_date = report_data.get(
        "end_date"
    )

    report_type = report_data.get(
        "report_type"
    )

    generated_at = _report_display_time(
        datetime.now(
            timezone.utc
        )
    )

    story.append(
        Paragraph(
            (
                f"{company_name} "
                "Media Intelligence Report"
            ),
            styles["title"],
        )
    )

    story.append(
        _pdf_paragraph(
            company_name,
            styles["subtitle"],
        )
    )

    if report_type:
        story.append(
            _pdf_paragraph(
                (
                    "Report Type: "
                    f"{str(report_type).title()}"
                ),
                styles["subtitle"],
            )
        )

    if (
        isinstance(start_date, datetime)
        and isinstance(end_date, datetime)
    ):
        display_start = _report_display_time(
            start_date
        )
        display_end = _report_display_time(
            end_date
        )

        period_text = (
            "Reporting Period: "
            f"{display_start:%d %b %Y %H:%M %Z}"
            " - "
            f"{display_end:%d %b %Y %H:%M %Z}"
        )
    else:
        period_text = (
            "Reporting Period: "
            "Not available"
        )

    story.append(
        _pdf_paragraph(
            period_text,
            styles["subtitle"],
        )
    )

    story.append(
        _pdf_paragraph(
            (
                "Generated: "
                f"{generated_at:%d %b %Y %H:%M %Z}"
            ),
            styles["small"],
        )
    )

    story.append(
        Spacer(
            1,
            0.28 * inch,
        )
    )

    story.append(
        Paragraph(
            "Executive Summary",
            styles["section"],
        )
    )

    story.append(
        _pdf_paragraph(
            report_data.get(
                "executive_summary"
            ),
            styles["body"],
            fallback=(
                "No executive summary "
                "is available."
            ),
        )
    )

    story.append(
        Spacer(
            1,
            0.14 * inch,
        )
    )

    story.append(
        Paragraph(
            "KPI Summary",
            styles["section"],
        )
    )

    sentiment_balance = report_data.get(
        "sentiment_balance",
        {},
    )
    risk_summary = report_data.get(
        "risk",
        {},
    )
    alerts = report_data.get(
        "alerts",
        [],
    )

    kpi_rows = [
        (
            "Total Articles",
            report_data.get(
                "total_articles",
                0,
            ),
        ),
        (
            "Total Events",
            report_data.get(
                "total_events",
                0,
            ),
        ),
        (
            "High Risk",
            report_data.get(
                "high_risk_count",
                0,
            ),
        ),
        (
            "Medium Risk",
            report_data.get(
                "medium_risk_count",
                0,
            ),
        ),
        (
            "Low Risk",
            report_data.get(
                "low_risk_count",
                0,
            ),
        ),
        (
            "Positive Sentiment",
            sentiment_balance.get(
                "positive",
                0,
            ),
        ),
        (
            "Neutral Sentiment",
            sentiment_balance.get(
                "neutral",
                0,
            ),
        ),
        (
            "Negative Sentiment",
            sentiment_balance.get(
                "negative",
                0,
            ),
        ),
        (
            "Alerts",
            len(alerts),
        ),
        (
            "Average Risk Score",
            risk_summary.get(
                "average_risk_score",
                0.0,
            ),
        ),
        (
            "Highest Risk Score",
            risk_summary.get(
                "highest_risk_score",
                0.0,
            ),
        ),
    ]

    kpi_table_data = [
        [
            Paragraph(
                "Metric",
                styles["table_header"],
            ),
            Paragraph(
                "Value",
                styles["table_header"],
            ),
        ]
    ]

    for label, value in kpi_rows:
        kpi_table_data.append(
            [
                _pdf_paragraph(
                    label,
                    styles["table_cell"],
                ),
                _pdf_paragraph(
                    _pdf_display_value(value),
                    styles["table_cell"],
                ),
            ]
        )

    kpi_table = Table(
        kpi_table_data,
        colWidths=[
            3.65 * inch,
            2.25 * inch,
        ],
        repeatRows=1,
        hAlign="LEFT",
    )

    kpi_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#16324F"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#D1D5DB"),
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F8FAFC"),
                    ],
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(
        kpi_table
    )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Sentiment",
            styles["section"],
        )
    )

    sentiment = report_data.get(
        "sentiment",
        {},
    )

    sentiment_rows = [
        (
            "Positive",
            sentiment.get(
                "positive",
                0,
            ),
        ),
        (
            "Neutral",
            sentiment.get(
                "neutral",
                0,
            ),
        ),
        (
            "Negative",
            sentiment.get(
                "negative",
                0,
            ),
        ),
    ]

    sentiment_table_data = [
        [
            Paragraph(
                "Classification",
                styles["table_header"],
            ),
            Paragraph(
                "Articles",
                styles["table_header"],
            ),
        ]
    ]

    for label, value in sentiment_rows:
        sentiment_table_data.append(
            [
                _pdf_paragraph(
                    label,
                    styles["table_cell"],
                ),
                _pdf_paragraph(
                    _pdf_display_value(value),
                    styles["table_cell"],
                ),
            ]
        )

    sentiment_table = Table(
        sentiment_table_data,
        colWidths=[
            3.65 * inch,
            2.25 * inch,
        ],
        repeatRows=1,
        hAlign="LEFT",
    )

    sentiment_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#16324F"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#D1D5DB"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F8FAFC"),
                    ],
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(
        sentiment_table
    )

    sentiment_series = sentiment.get(
        "series",
        [],
    )

    if sentiment_series:
        story.append(
            Spacer(
                1,
                0.10 * inch,
            )
        )

        sentiment_series_data = [
            [
                Paragraph(
                    "Period",
                    styles["table_header"],
                ),
                Paragraph(
                    "Positive",
                    styles["table_header"],
                ),
                Paragraph(
                    "Neutral",
                    styles["table_header"],
                ),
                Paragraph(
                    "Negative",
                    styles["table_header"],
                ),
            ]
        ]

        for item in sentiment_series:
            sentiment_series_data.append(
                [
                    _pdf_paragraph(
                        item.get("period"),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "positive",
                                0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "neutral",
                                0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "negative",
                                0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                ]
            )

        sentiment_series_table = Table(
            sentiment_series_data,
            colWidths=[
                2.3 * inch,
                1.2 * inch,
                1.2 * inch,
                1.2 * inch,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        sentiment_series_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#334E68"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.3,
                        colors.HexColor("#D1D5DB"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#F8FAFC"),
                        ],
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(
            sentiment_series_table
        )
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No sentiment trend data "
                    "is available for this period."
                ),
                styles["empty"],
            )
        )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Risk",
            styles["section"],
        )
    )

    risk = report_data.get(
        "risk",
        {},
    )

    risk_rows = [
        (
            "Average Risk Score",
            risk.get(
                "average_risk_score",
                0.0,
            ),
        ),
        (
            "Highest Risk Score",
            risk.get(
                "highest_risk_score",
                0.0,
            ),
        ),
        (
            "Low Risk",
            risk.get(
                "low_risk_count",
                0,
            ),
        ),
        (
            "Medium Risk",
            risk.get(
                "medium_risk_count",
                0,
            ),
        ),
        (
            "High Risk",
            risk.get(
                "high_risk_count",
                0,
            ),
        ),
    ]

    risk_table_data = [
        [
            Paragraph(
                "Metric",
                styles["table_header"],
            ),
            Paragraph(
                "Value",
                styles["table_header"],
            ),
        ]
    ]

    for label, value in risk_rows:
        risk_table_data.append(
            [
                _pdf_paragraph(
                    label,
                    styles["table_cell"],
                ),
                _pdf_paragraph(
                    _pdf_display_value(value),
                    styles["table_cell"],
                ),
            ]
        )

    risk_table = Table(
        risk_table_data,
        colWidths=[
            3.65 * inch,
            2.25 * inch,
        ],
        repeatRows=1,
        hAlign="LEFT",
    )

    risk_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#16324F"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#D1D5DB"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F8FAFC"),
                    ],
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(
        risk_table
    )

    risk_series = risk.get(
        "series",
        [],
    )

    if risk_series:
        story.append(
            Spacer(
                1,
                0.10 * inch,
            )
        )

        risk_series_data = [
            [
                Paragraph(
                    "Period",
                    styles["table_header"],
                ),
                Paragraph(
                    "Average",
                    styles["table_header"],
                ),
                Paragraph(
                    "Maximum",
                    styles["table_header"],
                ),
                Paragraph(
                    "Low",
                    styles["table_header"],
                ),
                Paragraph(
                    "Medium",
                    styles["table_header"],
                ),
                Paragraph(
                    "High",
                    styles["table_header"],
                ),
                Paragraph(
                    "Escalations",
                    styles["table_header"],
                ),
            ]
        ]

        for item in risk_series:
            risk_series_data.append(
                [
                    _pdf_paragraph(
                        item.get("period"),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "average_risk_score",
                                0.0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "maximum_risk_score",
                                0.0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "low_count",
                                0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "medium_count",
                                0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "high_count",
                                0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "escalation_count",
                                0,
                            )
                        ),
                        styles["table_cell"],
                    ),
                ]
            )

        risk_series_table = Table(
            risk_series_data,
            colWidths=[
                1.45 * inch,
                0.78 * inch,
                0.78 * inch,
                0.55 * inch,
                0.68 * inch,
                0.55 * inch,
                0.85 * inch,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        risk_series_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#334E68"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.3,
                        colors.HexColor("#D1D5DB"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#F8FAFC"),
                        ],
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(
            risk_series_table
        )
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No risk trend data is "
                    "available for this period."
                ),
                styles["empty"],
            )
        )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Business Impact",
            styles["section"],
        )
    )

    business_impact = report_data.get(
        "business_impact",
        {},
    )

    primary_distribution = business_impact.get(
        "primary_distribution",
        {},
    )
    category_distribution = business_impact.get(
        "category_distribution",
        {},
    )

    if any(primary_distribution.values()):
        primary_data = [
            [
                Paragraph(
                    "Primary Category",
                    styles["table_header"],
                ),
                Paragraph(
                    "Articles",
                    styles["table_header"],
                ),
            ]
        ]

        for category, count in primary_distribution.items():
            if count:
                primary_data.append(
                    [
                        _pdf_paragraph(
                            str(category).replace(
                                "_",
                                " ",
                            ).title(),
                            styles["table_cell"],
                        ),
                        _pdf_paragraph(
                            count,
                            styles["table_cell"],
                        ),
                    ]
                )

        primary_table = Table(
            primary_data,
            colWidths=[
                3.65 * inch,
                2.25 * inch,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        primary_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#16324F"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.35,
                        colors.HexColor("#D1D5DB"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#F8FAFC"),
                        ],
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        story.append(
            Paragraph(
                "Primary Distribution",
                styles["small"],
            )
        )
        story.append(primary_table)
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No primary business-impact "
                    "classifications are available "
                    "for this period."
                ),
                styles["empty"],
            )
        )

    if any(category_distribution.values()):
        story.append(
            Spacer(
                1,
                0.10 * inch,
            )
        )

        category_data = [
            [
                Paragraph(
                    "Impact Category",
                    styles["table_header"],
                ),
                Paragraph(
                    "Articles",
                    styles["table_header"],
                ),
            ]
        ]

        for category, count in category_distribution.items():
            if count:
                category_data.append(
                    [
                        _pdf_paragraph(
                            str(category).replace(
                                "_",
                                " ",
                            ).title(),
                            styles["table_cell"],
                        ),
                        _pdf_paragraph(
                            count,
                            styles["table_cell"],
                        ),
                    ]
                )

        category_table = Table(
            category_data,
            colWidths=[
                3.65 * inch,
                2.25 * inch,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        category_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#334E68"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.35,
                        colors.HexColor("#D1D5DB"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#F8FAFC"),
                        ],
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        story.append(
            Paragraph(
                "All-Category Distribution",
                styles["small"],
            )
        )
        story.append(category_table)
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No all-category business-impact "
                    "classifications are available "
                    "for this period."
                ),
                styles["empty"],
            )
        )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Events",
            styles["section"],
        )
    )

    events = report_data.get(
        "events",
        {},
    )
    event_rows = events.get(
        "largest_events",
        [],
    )

    if event_rows:
        event_table_data = [
            [
                Paragraph(
                    "Event",
                    styles["table_header"],
                ),
                Paragraph(
                    "Articles",
                    styles["table_header"],
                ),
                Paragraph(
                    "First Published",
                    styles["table_header"],
                ),
                Paragraph(
                    "Last Published",
                    styles["table_header"],
                ),
            ]
        ]

        for item in event_rows:
            event_table_data.append(
                [
                    _pdf_paragraph(
                        item.get("title"),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        item.get(
                            "article_count",
                            0,
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "first_published_at"
                            )
                        ),
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            item.get(
                                "last_published_at"
                            )
                        ),
                        styles["table_cell"],
                    ),
                ]
            )

        event_table = Table(
            event_table_data,
            colWidths=[
                2.6 * inch,
                0.65 * inch,
                1.35 * inch,
                1.35 * inch,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        event_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#16324F"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.3,
                        colors.HexColor("#D1D5DB"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#F8FAFC"),
                        ],
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(event_table)
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No qualifying event clusters "
                    "were identified for this period."
                ),
                styles["empty"],
            )
        )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Highest Risk Stories",
            styles["section"],
        )
    )

    highest_risk_stories = report_data.get(
        "highest_risk_stories",
        [],
    )

    if highest_risk_stories:
        for index, item in enumerate(
            highest_risk_stories,
            start=1,
        ):
            title = item.get(
                "title",
                "Untitled story",
            )

            source_name = item.get(
                "publisher_name"
            ) or item.get("source_name")
            published_at = item.get(
                "published_at"
            )
            collected_at = item.get(
                "collected_at"
            )
            risk_score = item.get(
                "risk_score"
            )
            risk_level = item.get(
                "risk_level"
            )
            sentiment_value = item.get(
                "sentiment"
            )
            business_impact = item.get(
                "business_impact_primary"
            )
            risk_headline = item.get(
                "risk_headline"
            )
            risk_summary = item.get(
                "risk_summary"
            )
            event_title = item.get(
                "event_cluster_title"
            )

            metadata_parts = []

            if source_name:
                metadata_parts.append(
                    f"Source: {source_name}"
                )

            if published_at:
                metadata_parts.append(
                    (
                        "Published: "
                        f"{_pdf_display_value(published_at)}"
                    )
                )

            if collected_at:
                metadata_parts.append(
                    (
                        "Added to Nova Cops: "
                        f"{_pdf_display_value(collected_at)}"
                    )
                )

            if risk_score is not None:
                metadata_parts.append(
                    (
                        "Risk Score: "
                        f"{_pdf_display_value(risk_score)}"
                    )
                )

            if risk_level:
                metadata_parts.append(
                    (
                        "Risk Level: "
                        f"{str(risk_level).title()}"
                    )
                )

            if sentiment_value:
                metadata_parts.append(
                    (
                        "Sentiment: "
                        f"{str(sentiment_value).title()}"
                    )
                )

            if business_impact:
                metadata_parts.append(
                    (
                        "Business Impact: "
                        f"{str(business_impact).replace('_', ' ').title()}"
                    )
                )

            story_block = [
                _pdf_paragraph(
                    f"{index}. {title}",
                    styles["body"],
                )
            ]

            if metadata_parts:
                story_block.append(
                    _pdf_paragraph(
                        " | ".join(
                            metadata_parts
                        ),
                        styles["small"],
                    )
                )

            if risk_headline:
                story_block.append(
                    _pdf_paragraph(
                        (
                            "Risk Headline: "
                            f"{risk_headline}"
                        ),
                        styles["body"],
                    )
                )

            if risk_summary:
                story_block.append(
                    _pdf_paragraph(
                        risk_summary,
                        styles["body"],
                    )
                )

            if event_title:
                story_block.append(
                    _pdf_paragraph(
                        (
                            "Event Cluster: "
                            f"{event_title}"
                        ),
                        styles["small"],
                    )
                )

            story_block.append(
                Spacer(
                    1,
                    0.10 * inch,
                )
            )

            story.append(
                KeepTogether(
                    story_block
                )
            )
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No risk-ranked stories are "
                    "available for this period."
                ),
                styles["empty"],
            )
        )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Sources",
            styles["section"],
        )
    )

    sources = report_data.get(
        "sources",
        [],
    )

    if sources:
        source_table_data = [
            [
                Paragraph(
                    "Source",
                    styles["table_header"],
                ),
                Paragraph(
                    "Articles",
                    styles["table_header"],
                ),
            ]
        ]

        for item in sources:
            source_name = (
                item.get("source_name")
                or item.get("source")
                or item.get("name")
                or "Unknown source"
            )

            article_count = (
                item.get("article_count")
                if item.get("article_count") is not None
                else item.get("count", 0)
            )

            source_table_data.append(
                [
                    _pdf_paragraph(
                        source_name,
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            article_count
                        ),
                        styles["table_cell"],
                    ),
                ]
            )

        source_table = Table(
            source_table_data,
            colWidths=[
                4.25 * inch,
                1.65 * inch,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        source_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#16324F"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.35,
                        colors.HexColor("#D1D5DB"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#F8FAFC"),
                        ],
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        story.append(source_table)
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No qualifying media sources "
                    "were identified for this period."
                ),
                styles["empty"],
            )
        )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Competitors",
            styles["section"],
        )
    )

    competitors = report_data.get(
        "competitors",
        [],
    )

    if competitors:
        competitor_table_data = [
            [
                Paragraph(
                    "Competitor",
                    styles["table_header"],
                ),
                Paragraph(
                    "Mentions",
                    styles["table_header"],
                ),
            ]
        ]

        for item in competitors:
            competitor_name = (
                item.get("competitor_name")
                or item.get("company_name")
                or item.get("name")
                or "Configured competitor"
            )

            mention_count = item.get(
                "mention_count",
                0,
            )

            competitor_table_data.append(
                [
                    _pdf_paragraph(
                        competitor_name,
                        styles["table_cell"],
                    ),
                    _pdf_paragraph(
                        _pdf_display_value(
                            mention_count
                        ),
                        styles["table_cell"],
                    ),
                ]
            )

        competitor_table = Table(
            competitor_table_data,
            colWidths=[
                4.25 * inch,
                1.65 * inch,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        competitor_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#16324F"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.35,
                        colors.HexColor("#D1D5DB"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#F8FAFC"),
                        ],
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        story.append(
            competitor_table
        )
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No competitors are currently "
                    "configured for comparison."
                ),
                styles["empty"],
            )
        )

    story.append(
        Spacer(
            1,
            0.16 * inch,
        )
    )

    story.append(
        Paragraph(
            "Alerts",
            styles["section"],
        )
    )

    alert_rows = report_data.get(
        "alerts",
        [],
    )

    if alert_rows:
        for index, item in enumerate(
            alert_rows,
            start=1,
        ):
            title = item.get(
                "title",
                "Alert",
            )

            severity = item.get(
                "severity"
            )
            alert_type = item.get(
                "alert_type"
            )
            message = item.get(
                "message"
            )
            delivery_status = item.get(
                "delivery_status"
            )
            delivery_channel = item.get(
                "delivery_channel"
            )
            created_at = item.get(
                "created_at"
            )
            delivered_at = item.get(
                "delivered_at"
            )
            sla_due_at = item.get(
                "sla_due_at"
            )

            metadata_parts = []

            if severity:
                metadata_parts.append(
                    (
                        "Severity: "
                        f"{str(severity).title()}"
                    )
                )

            if alert_type:
                metadata_parts.append(
                    (
                        "Type: "
                        f"{str(alert_type).replace('_', ' ').title()}"
                    )
                )

            if delivery_status:
                metadata_parts.append(
                    (
                        "Status: "
                        f"{str(delivery_status).title()}"
                    )
                )

            if delivery_channel:
                metadata_parts.append(
                    (
                        "Channel: "
                        f"{str(delivery_channel).title()}"
                    )
                )

            alert_block = [
                _pdf_paragraph(
                    f"{index}. {title}",
                    styles["body"],
                )
            ]

            if metadata_parts:
                alert_block.append(
                    _pdf_paragraph(
                        " | ".join(
                            metadata_parts
                        ),
                        styles["small"],
                    )
                )

            if message:
                alert_block.append(
                    _pdf_paragraph(
                        message,
                        styles["body"],
                    )
                )

            time_parts = []

            if created_at:
                time_parts.append(
                    (
                        "Created: "
                        f"{_pdf_display_value(created_at)}"
                    )
                )

            if sla_due_at:
                time_parts.append(
                    (
                        "SLA Due: "
                        f"{_pdf_display_value(sla_due_at)}"
                    )
                )

            if delivered_at:
                time_parts.append(
                    (
                        "Delivered: "
                        f"{_pdf_display_value(delivered_at)}"
                    )
                )

            if time_parts:
                alert_block.append(
                    _pdf_paragraph(
                        " | ".join(time_parts),
                        styles["small"],
                    )
                )

            alert_block.append(
                Spacer(
                    1,
                    0.08 * inch,
                )
            )

            story.append(
                KeepTogether(
                    alert_block
                )
            )
    else:
        story.append(
            _pdf_paragraph(
                (
                    "No alerts were generated "
                    "during the reporting period."
                ),
                styles["empty"],
            )
        )

    document.build(
        story,
        onFirstPage=_pdf_page_footer,
        onLaterPages=_pdf_page_footer,
    )

    return output.getvalue()

def render_report(
    report_data: dict,
    file_format: str,
) -> bytes:
    if file_format not in REPORT_FORMATS:
        raise ValueError(
            (
                "Unsupported report format: "
                f"{file_format}"
            )
        )

    if _is_article_collection_report(report_data):
        exporters = {
            "pdf": export_article_collection_pdf,
            "xlsx": export_article_collection_xlsx,
            "csv": export_article_collection_csv,
        }
    else:
        exporters = {
        "pdf": export_report_pdf,
        "xlsx": export_report_xlsx,
        "csv": export_report_csv,
        }

    return exporters[
        file_format
    ](
        report_data
    )


async def find_existing_report(
    db: AsyncSession,
    *,
    company_id: int,
    report_type: str,
    file_format: str,
    period_start: datetime,
    period_end: datetime,
) -> GeneratedReport | None:
    stmt = select(
        GeneratedReport
    ).where(
        GeneratedReport.company_id
        == company_id,
        GeneratedReport.report_type
        == report_type,
        GeneratedReport.file_format
        == file_format,
        GeneratedReport.period_start
        == period_start,
        GeneratedReport.period_end
        == period_end,
    )

    return (
        await db.execute(
            stmt
        )
    ).scalar_one_or_none()


async def create_pending_report(
    db: AsyncSession,
    *,
    company_id: int,
    report_type: str,
    file_format: str,
    period_start: datetime,
    period_end: datetime,
) -> GeneratedReport:
    if file_format not in REPORT_FORMATS:
        raise ValueError(
            (
                "Unsupported report format: "
                f"{file_format}"
            )
        )

    period_start, period_end = (
        validate_time_window(
            period_start,
            period_end,
        )
    )

    existing = await find_existing_report(
        db,
        company_id=company_id,
        report_type=report_type,
        file_format=file_format,
        period_start=period_start,
        period_end=period_end,
    )

    if existing is not None:
        return existing

    record = GeneratedReport(
        company_id=company_id,
        report_type=report_type,
        file_format=file_format,
        period_start=period_start,
        period_end=period_end,
        filename=None,
        content_type=None,
        content=None,
        status="pending",
        generated_at=None,
        error=None,
    )

    db.add(record)
    await db.commit()
    await db.refresh(record)

    return record


async def mark_report_processing(
    db: AsyncSession,
    record: GeneratedReport,
) -> GeneratedReport:
    record.status = "processing"
    record.error = None

    await db.commit()
    await db.refresh(record)

    return record


async def mark_report_success(
    db: AsyncSession,
    record: GeneratedReport,
    *,
    content: bytes,
) -> GeneratedReport:
    extension = REPORT_EXTENSIONS[
        record.file_format
    ]

    record.filename = (
        f"company-{record.company_id}-"
        f"{record.report_type}-"
        f"{record.period_start.date().isoformat()}-"
        f"{record.period_end.date().isoformat()}."
        f"{extension}"
    )

    record.content_type = (
        REPORT_CONTENT_TYPES[
            record.file_format
        ]
    )

    record.content = content
    record.status = "success"
    record.generated_at = datetime.now(
        timezone.utc
    )
    record.error = None

    await db.commit()
    await db.refresh(record)

    return record


async def mark_report_failed(
    db: AsyncSession,
    record: GeneratedReport,
    *,
    error: str,
) -> GeneratedReport:
    record.status = "failed"
    record.error = error[:1000]
    record.generated_at = None
    record.filename = None
    record.content_type = None
    record.content = None

    await db.commit()
    await db.refresh(record)

    return record


async def generate_report_record(
    db: AsyncSession,
    *,
    company_id: int,
    report_type: str,
    file_format: str,
    period_start: datetime,
    period_end: datetime,
) -> GeneratedReport:
    record = await create_pending_report(
        db,
        company_id=company_id,
        report_type=report_type,
        file_format=file_format,
        period_start=period_start,
        period_end=period_end,
    )

    if record.status == "success":
        return record

    await mark_report_processing(
        db,
        record,
    )

    try:
        report_data = await build_company_report(
            db,
            company_id=company_id,
            start_date=period_start,
            end_date=period_end,
        )

        content = render_report(
            report_data,
            file_format,
        )

        return await mark_report_success(
            db,
            record,
            content=content,
        )

    except Exception as exc:
        await mark_report_failed(
            db,
            record,
            error=str(exc),
        )

        raise


async def persist_generated_report(
    db: AsyncSession,
    *,
    report_data: dict,
    file_format: str,
    content: bytes,
    report_type: str = "custom",
) -> GeneratedReport:
    """
    Compatibility wrapper for the current Phase 17 API/tasks.

    Phase 17G will migrate callers to generate_report_record().
    """

    record = await create_pending_report(
        db,
        company_id=report_data[
            "company_id"
        ],
        report_type=report_type,
        file_format=file_format,
        period_start=report_data[
            "start_date"
        ],
        period_end=report_data[
            "end_date"
        ],
    )

    if record.status == "success":
        return record

    await mark_report_processing(
        db,
        record,
    )

    return await mark_report_success(
        db,
        record,
        content=content,
    )


import uuid

REPORT_FORMATS = ("pdf", "xlsx", "csv")


async def generate_report_batch(
    db: AsyncSession,
    *,
    company_id: int,
    report_type: str,
    period_start: datetime,
    period_end: datetime,
    time_mode: str = "media",
    article_ids: list[int] | None = None,
    report_scope: str = "standard",
    report_title: str | None = None,
) -> list[GeneratedReport]:
    """
    Generate all three report formats (PDF, XLSX, CSV) in a single batch
    sharing one snapshot_at timestamp and one report_data build.
    Returns the list of GeneratedReport records (one per format).
    """
    snapshot_at = datetime.now(timezone.utc)
    batch_id = str(uuid.uuid4())

    records = []
    for file_format in REPORT_FORMATS:
        record = GeneratedReport(
            company_id=company_id,
            report_type=report_type,
            file_format=file_format,
            period_start=period_start,
            period_end=period_end,
            batch_id=batch_id,
            snapshot_at=snapshot_at,
            time_mode=time_mode,
            filename=None,
            content_type=None,
            content=None,
            status="processing",
            generated_at=None,
            error=None,
        )
        db.add(record)
        records.append(record)

    await db.commit()
    for r in records:
        await db.refresh(r)

    try:
        report_data = await build_company_report(
            db,
            company_id=company_id,
            start_date=period_start,
            end_date=period_end,
            snapshot_at=snapshot_at,
            time_mode=time_mode,
            article_ids=article_ids,
            report_scope=report_scope,
            report_title=report_title,
        )

        for record in records:
            content = render_report(report_data, record.file_format)
            extension = REPORT_EXTENSIONS[record.file_format]
            record.filename = (
                f"company-{record.company_id}-"
                f"{record.report_type}-"
                f"{record.period_start.date().isoformat()}-"
                f"{record.period_end.date().isoformat()}."
                f"{extension}"
            )
            record.content_type = REPORT_CONTENT_TYPES[record.file_format]
            record.content = content
            record.status = "success"
            record.generated_at = snapshot_at
            record.included_article_ids = [
                row["article_id"] for row in report_data.get("articles", [])
            ]
            record.error = None

        await db.commit()
        for r in records:
            await db.refresh(r)

        return records

    except Exception as exc:
        for record in records:
            record.status = "failed"
            record.error = str(exc)[:1000]
            record.generated_at = None
            record.filename = None
            record.content_type = None
            record.content = None

        await db.commit()
        raise
