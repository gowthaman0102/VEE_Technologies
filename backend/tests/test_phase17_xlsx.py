from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.services.report_service import (
    build_deterministic_executive_summary,
    export_report_xlsx,
)


EXPECTED_SHEETS = [
    "Summary",
    "Articles",
    "Sentiment",
    "Risk",
    "Business Impact",
    "Events",
    "Sources",
    "Competitors",
    "Alerts",
]


def _empty_report() -> dict:
    report = {
        "company_id": 2,
        "company_name": "VEE Technologies",
        "start_date": datetime(
            2026, 9, 1,
            tzinfo=timezone.utc,
        ),
        "end_date": datetime(
            2026, 9, 17,
            tzinfo=timezone.utc,
        ),
        "total_articles": 0,
        "total_events": 0,
        "high_risk_count": 0,
        "medium_risk_count": 0,
        "low_risk_count": 0,
        "sentiment_balance": {
            "positive": 0,
            "neutral": 0,
            "negative": 0,
        },
        "metrics": [],
        "overview": {
            "total_articles": 0,
            "total_events": 0,
            "total_alerts": 0,
        },
        "article_trend": [],
        "articles": [],
        "sentiment": {
            "positive": 0,
            "neutral": 0,
            "negative": 0,
            "series": [],
        },
        "risk": {
            "average_risk_score": 0.0,
            "highest_risk_score": 0.0,
            "high_risk_count": 0,
            "medium_risk_count": 0,
            "low_risk_count": 0,
            "series": [],
        },
        "business_impact": {
            "primary_distribution": {},
            "category_distribution": {},
        },
        "events": {
            "total_events": 0,
            "largest_events": [],
        },
        "sources": [],
        "competitors": [],
        "alerts": [],
        "highest_risk_stories": [],
    }

    report["executive_summary"] = (
        build_deterministic_executive_summary(
            report
        )
    )

    return report


def test_phase17_xlsx_empty_report_renders() -> None:
    report = _empty_report()

    content = export_report_xlsx(report)

    assert content.startswith(b"PK")
    assert len(content) > 8000

    workbook = load_workbook(
        BytesIO(content)
    )

    assert workbook.sheetnames == EXPECTED_SHEETS

    assert (
        workbook["Summary"]["A2"].value
        == "VEE Technologies"
    )

    assert (
        workbook["Articles"]["A2"].value
        == (
            "No qualifying articles are available "
            "for this reporting period."
        )
    )

    assert (
        workbook["Competitors"]["A2"].value
        == (
            "No competitors are currently "
            "configured for comparison."
        )
    )

    assert (
        workbook["Alerts"]["A2"].value
        == (
            "No alerts were generated "
            "during the reporting period."
        )
    )


def test_phase17_xlsx_nonempty_report_renders_real_rows() -> None:
    report = _empty_report()

    report.update(
        {
            "total_articles": 1,
            "total_events": 1,
            "high_risk_count": 1,
            "low_risk_count": 0,
        }
    )

    report["sentiment_balance"] = {
        "positive": 0,
        "neutral": 0,
        "negative": 1,
    }

    report["articles"] = [
        {
            "article_id": 101,
            "title": "Operational risk story",
            "source_name": "Business Daily",
            "published_at": (
                "2026-09-15T09:00:00+00:00"
            ),
            "url": "https://example.com/story",
            "sentiment": "negative",
            "sentiment_score": -0.8,
            "business_impact_primary": (
                "operational"
            ),
            "business_impact_categories": [
                "operational",
                "reputation",
            ],
            "business_impact_summary": (
                "Stored business impact."
            ),
            "competitors": [],
            "risk_score": 88.0,
            "risk_level": "high",
            "escalation_action": "escalate",
            "attention_level": "immediate",
            "risk_headline": (
                "Elevated operational risk"
            ),
            "risk_summary": (
                "Stored risk summary."
            ),
            "event_cluster_title": (
                "Operational event"
            ),
            "alerts": [
                {"title": "Risk alert"}
            ],
        }
    ]

    report["sentiment"] = {
        "positive": 0,
        "neutral": 0,
        "negative": 1,
        "series": [
            {
                "period": "2026-09-15",
                "positive": 0,
                "neutral": 0,
                "negative": 1,
            }
        ],
    }

    report["risk"] = {
        "average_risk_score": 88.0,
        "highest_risk_score": 88.0,
        "low_risk_count": 0,
        "medium_risk_count": 0,
        "high_risk_count": 1,
        "series": [
            {
                "period": "2026-09-15",
                "average_risk_score": 88.0,
                "maximum_risk_score": 88.0,
                "low_count": 0,
                "medium_count": 0,
                "high_count": 1,
                "escalation_count": 1,
            }
        ],
    }

    report["business_impact"] = {
        "primary_distribution": {
            "operational": 1,
        },
        "category_distribution": {
            "operational": 1,
            "reputation": 1,
        },
    }

    report["events"] = {
        "total_events": 1,
        "largest_events": [
            {
                "cluster_id": 10,
                "title": "Operational event",
                "article_count": 1,
                "first_published_at": (
                    "2026-09-15T09:00:00+00:00"
                ),
                "last_published_at": (
                    "2026-09-15T09:00:00+00:00"
                ),
            }
        ],
    }

    report["sources"] = [
        {
            "source_name": "Business Daily",
            "article_count": 1,
        }
    ]

    report["alerts"] = [
        {
            "article_id": 101,
            "alert_type": "risk_threshold",
            "severity": "high",
            "title": "Risk alert",
            "message": "Risk threshold exceeded.",
            "delivery_status": "delivered",
            "delivery_channel": "slack",
            "created_at": (
                "2026-09-15T09:05:00+00:00"
            ),
        }
    ]

    report["executive_summary"] = (
        build_deterministic_executive_summary(
            report
        )
    )

    content = export_report_xlsx(report)

    workbook = load_workbook(
        BytesIO(content)
    )

    assert (
        workbook["Articles"]["A2"].value
        == 101
    )

    assert (
        workbook["Articles"]["L2"].value
        == 88.0
    )

    assert (
        workbook["Sentiment"]["G3"].value
        == 1
    )

    assert (
        workbook["Risk"]["J3"].value
        == 1
    )

    assert (
        workbook["Business Impact"]["A2"].value
        == "Operational"
    )

    assert (
        workbook["Events"]["A2"].value
        == 10
    )

    assert (
        workbook["Sources"]["A2"].value
        == "Business Daily"
    )

    assert (
        workbook["Alerts"]["C2"].value
        == "high"
    )


def test_phase17_xlsx_has_no_phase17e_placeholders() -> None:
    report = _empty_report()

    content = export_report_xlsx(report)

    workbook = load_workbook(
        BytesIO(content)
    )

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert not (
                    isinstance(
                        cell.value,
                        str,
                    )
                    and (
                        "populated during Phase 17E"
                        in cell.value
                    )
                )


def test_phase17_xlsx_does_not_mutate_input() -> None:
    report = _empty_report()
    before = deepcopy(report)

    export_report_xlsx(report)

    assert report == before
