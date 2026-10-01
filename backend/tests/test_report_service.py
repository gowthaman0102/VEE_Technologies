from datetime import datetime, timezone
import csv
import io

from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.report_service import (
    export_report_csv,
    export_report_pdf,
    export_report_xlsx,
    render_article_list_pdf,
    render_report,
)
from app.services.report_schedule_service import calculate_next_run_at
from app.schemas.report_management import ReportRecipientsUpdate, ReportSchedulePayload
from pydantic import ValidationError
import pytest


REPORT = {
    "company_id": 1,
    "start_date": datetime(2026, 1, 1, tzinfo=timezone.utc),
    "end_date": datetime(2026, 1, 31, tzinfo=timezone.utc),
    "total_articles": 3,
    "total_events": 1,
    "high_risk_count": 1,
    "medium_risk_count": 1,
    "low_risk_count": 1,
    "sentiment_balance": {"positive": 1, "neutral": 1, "negative": 1},
    "metrics": [{"label": "Average risk score", "value": 55.0}],
}


def test_report_renderers_produce_real_files():
    csv_content = export_report_csv(REPORT)
    xlsx_content = export_report_xlsx(REPORT)
    pdf_content = export_report_pdf(REPORT)

    assert csv_content.startswith(b"row_type,metric_label")
    assert xlsx_content.startswith(b"PK")
    assert pdf_content.startswith(b"%PDF")


def test_summary_template_omits_article_details_but_keeps_aggregates():
    report_data = {
        **REPORT,
        "articles": [{"article_id": 17, "title": "Detailed article title"}],
        "highest_risk_stories": [{"article_id": 17, "title": "Detailed article title"}],
        "alerts": [{"id": 3, "title": "Detailed alert"}],
        "competitors": [{"name": "Detailed competitor"}],
    }

    detailed = render_report(report_data, "csv", include_details=True)
    summary = render_report(report_data, "csv", include_details=False)

    assert b"Detailed article title" in detailed
    assert b"Detailed article title" not in summary
    assert b"total_articles,3" in summary
    assert report_data["articles"][0]["article_id"] == 17


def test_standard_report_templates_have_distinct_compact_content():
    report_data = {
        **REPORT,
        "company_name": "Example Co",
        "time_mode": "media",
        "critical_risk_count": 2,
        "risk": {"average_risk_score": 68.5, "highest_risk_score": 91.0},
        "business_impact": {
            "category_distribution": {"cybersecurity": 6, "legal": 3},
        },
        "executive_summary": "Stored, deterministic summary.",
        "metrics": [{"label": "Processed Intelligence", "value": 8}],
        "articles": [
            {"article_id": article_id, "title": f"Full appendix article {article_id}"}
            for article_id in range(1, 11)
        ],
        "highest_risk_stories": [
            {
                "article_id": article_id,
                "title": f"Priority story {article_id}",
                "risk_level": "critical",
                "risk_score": 90 - article_id,
            }
            for article_id in range(1, 11)
        ],
        "alerts": [{"id": 1, "title": "Stored alert"}],
    }

    brief = render_report(report_data, "csv", report_template="executive")
    board = render_report(report_data, "csv", report_template="board_ready")
    full = render_report(report_data, "csv", report_template="detailed")
    brief_text = brief.decode("utf-8")
    board_text = board.decode("utf-8")
    full_text = full.decode("utf-8")

    assert "Brief Summary" in brief_text
    assert "Board Summary" in board_text
    assert "board_signal" in board_text
    assert "Full appendix article" not in brief_text
    assert "Full appendix article" not in board_text
    assert "Full appendix article 10" in full_text
    assert "Priority story 8" in brief_text
    assert "Priority story 9" not in brief_text
    assert "Priority story 5" in board_text
    assert "Priority story 6" not in board_text

    brief_xlsx = load_workbook(
        io.BytesIO(render_report(report_data, "xlsx", report_template="executive")),
        read_only=True,
    )
    board_xlsx = load_workbook(
        io.BytesIO(render_report(report_data, "xlsx", report_template="board_ready")),
        read_only=True,
    )
    assert brief_xlsx.sheetnames == ["Summary", "Material Stories"]
    assert board_xlsx.sheetnames == ["Summary", "Material Stories"]
    assert render_report(report_data, "pdf", report_template="executive") != render_report(
        report_data,
        "pdf",
        report_template="board_ready",
    )


def test_article_pdf_preserves_japanese_headline_and_em_dash():
    pdf = render_article_list_pdf({
        "company_name": "Example Co",
        "report_scope": "search_results",
        "search_query": "OpenAI",
        "search_mode": "keyword",
        "articles": [
            {
                "article_id": 1,
                "title": "パトロンプラットフォーム Patreon — full headline",
                "publisher_name": "Example News",
                "url": "https://example.com/article",
            },
            {
                "article_id": 2,
                "title": "The Era of Token Abundance Is Coming — full headline",
                "publisher_name": "Example News",
                "url": "https://example.com/article-2",
            },
            {
                "article_id": 3,
                "title": "The Era of Token Abundance Is Coming " + bytes((0xE2, 0x80, 0x94)).decode("latin-1") + " mojibake headline",
                "publisher_name": "Example News",
                "url": "https://example.com/article-3",
            },
        ],
    })
    text = " ".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)

    assert "パトロンプラットフォーム Patreon" in text
    assert "Patreon - full headline" in text
    assert "The Era of Token Abundance Is Coming - full headline" in text
    assert "The Era of Token Abundance Is Coming - mojibake headline" in text
    assert "The Era of Token Abundance Is Coming - full headline" in text


def test_report_schedule_next_run_uses_utc_cadence_fields():
    now = datetime(2026, 9, 28, 10, tzinfo=timezone.utc)

    assert calculate_next_run_at(
        report_type="daily", run_time_utc="09:00", now=now
    ) == datetime(2026, 9, 29, 9, tzinfo=timezone.utc)
    assert calculate_next_run_at(
        report_type="weekly", run_time_utc="11:00", day_of_week=0, now=now
    ) == datetime(2026, 9, 28, 11, tzinfo=timezone.utc)
    assert calculate_next_run_at(
        report_type="monthly", run_time_utc="09:00", day_of_month=1, now=now
    ) == datetime(2026, 10, 1, 9, tzinfo=timezone.utc)


def test_report_management_schemas_validate_live_inputs():
    recipients = ReportRecipientsUpdate(emails=["Person@Example.com", " person@example.com "])
    assert recipients.emails == ["person@example.com"]

    with pytest.raises(ValidationError):
        ReportRecipientsUpdate(emails=["not-an-email"])
    with pytest.raises(ValidationError):
        ReportSchedulePayload(name="Impact", report_type="daily", report_scope="business_impact")
